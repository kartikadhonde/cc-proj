import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def generate_team_guide_docx():
    doc = Document()

    # 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base typography
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    input_path = os.path.join('docs', 'TEAM_CONTRIBUTION_AND_PRESENTATION_GUIDE.md')
    with open(input_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_text = []
    table_lines = []

    def flush_table(t_lines):
        if not t_lines or len(t_lines) < 2:
            return
        rows = []
        for l in t_lines:
            parts = [p.strip() for p in l.strip().strip('|').split('|')]
            rows.append(parts)
        cleaned_rows = [r for r in rows if not all(set(c) <= set('-: ') for c in r)]
        if not cleaned_rows:
            return
        num_cols = max(len(r) for r in cleaned_rows)
        t = doc.add_table(rows=len(cleaned_rows), cols=num_cols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r_idx, row_data in enumerate(cleaned_rows):
            for c_idx, val in enumerate(row_data):
                if c_idx < num_cols:
                    cell = t.cell(r_idx, c_idx)
                    cleaned_val = (val.replace('**', '')
                                      .replace('`', ''))
                    cell.text = cleaned_val
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    for run in p.runs:
                        run.font.size = Pt(9.5)
                        if r_idx == 0:
                            run.font.bold = True
                    if r_idx == 0:
                        shd_xml = f'<w:shd {nsdecls("w")} w:fill="EBF2F7"/>'
                        cell._tc.get_or_add_tcPr().append(parse_xml(shd_xml))
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        
        # Code blocks
        if line.startswith('```'):
            if in_code_block:
                in_code_block = False
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                run = p.add_run('\n'.join(code_text))
                run.font.name = 'Consolas'
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
                code_text = []
            else:
                in_code_block = True
                code_text = []
            i += 1
            continue

        if in_code_block:
            code_text.append(line)
            i += 1
            continue

        # Markdown Tables
        if line.startswith('|') and '|' in line[1:]:
            table_lines.append(line)
            i += 1
            continue
        elif table_lines:
            flush_table(table_lines)
            table_lines = []

        # Empty lines
        if not line.strip():
            i += 1
            continue

        # Horizontal rules
        if line.strip() in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CCCCCC"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            i += 1
            continue

        # Headings
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[2:].replace('**', '').replace('`', ''))
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D) # Navy
            i += 1
            continue

        if line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[3:].replace('**', '').replace('`', ''))
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x23, 0x5A, 0x8C) # Slate Blue
            i += 1
            continue

        if line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[4:].replace('**', '').replace('`', ''))
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            i += 1
            continue

        if line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[5:].replace('**', '').replace('`', ''))
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            i += 1
            continue

        # Blockquotes (Speaking scripts)
        if line.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(4)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="8" w:color="235A8C"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            quote_text = line[2:]
            _format_inline_text(p, quote_text, italic=True)
            i += 1
            continue

        # Bullet lists
        if line.lstrip().startswith(('* ', '- ', '• ')):
            indent_level = (len(line) - len(line.lstrip())) // 2
            text = line.lstrip()[2:]
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent = Inches(0.25 * (indent_level + 1))
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            _format_inline_text(p, text)
            i += 1
            continue

        # Numbered lists
        stripped = line.lstrip()
        if len(stripped) > 2 and stripped[0].isdigit() and stripped[1:3] in ('. ', ') '):
            indent_level = (len(line) - len(stripped)) // 2
            text = stripped[3:]
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.left_indent = Inches(0.25 * (indent_level + 1))
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            _format_inline_text(p, text)
            i += 1
            continue

        # Regular paragraphs
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        _format_inline_text(p, line)
        i += 1

    if table_lines:
        flush_table(table_lines)

    output_path = os.path.join('docs', 'TEAM_CONTRIBUTION_AND_PRESENTATION_GUIDE.docx')
    doc.save(output_path)
    print(f"Successfully generated: {output_path}")

def _format_inline_text(paragraph, text, italic=False):
    parts = text.split('**')
    for idx, part in enumerate(parts):
        if not part:
            continue
        is_bold = (idx % 2 == 1)
        sub_parts = part.split('`')
        for s_idx, s_part in enumerate(sub_parts):
            if not s_part:
                continue
            is_code = (s_idx % 2 == 1)
            run = paragraph.add_run(s_part)
            if is_bold:
                run.bold = True
            if italic:
                run.italic = True
            if is_code:
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x99, 0x22, 0x22)

if __name__ == '__main__':
    generate_team_guide_docx()
