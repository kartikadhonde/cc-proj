import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def generate_report_docx():
    doc = Document()

    # Set 1-inch standard margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base typography styling
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    with open('PROJECT_REPORT.md', 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_text = []
    table_lines = []

    def flush_table(t_lines):
        if not t_lines or len(t_lines) < 2:
            return
        rows = []
        for l in t_lines:
            parts = [p.strip() for p in l.strip().strip('|').split('|')]
            rows.append(parts)
        cleaned_rows = [r for r in rows if not all(set(c) <= set('-: ') for c in r)]
        if not cleaned_rows:
            return
        num_cols = max(len(r) for r in cleaned_rows)
        t = doc.add_table(rows=len(cleaned_rows), cols=num_cols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r_idx, row_data in enumerate(cleaned_rows):
            for c_idx, val in enumerate(row_data):
                if c_idx < num_cols:
                    cell = t.cell(r_idx, c_idx)
                    cleaned_val = (val.replace('**', '')
                                      .replace('`', '')
                                      .replace('$', '')
                                      .replace('\\text{', '')
                                      .replace('}', ''))
                    cell.text = cleaned_val
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    for run in p.runs:
                        run.font.size = Pt(9.5)
                        if r_idx == 0:
                            run.font.bold = True
                    if r_idx == 0:
                        shd_xml = f'<w:shd {nsdecls("w")} w:fill="EBF2F7"/>'
                        cell._tc.get_or_add_tcPr().append(parse_xml(shd_xml))
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        
        # Code blocks
        if line.startswith('```'):
            if in_code_block:
                in_code_block = False
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                run = p.add_run('\n'.join(code_text))
                run.font.name = 'Consolas'
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
                code_text = []
            else:
                in_code_block = True
                code_text = []
            i += 1
            continue

        if in_code_block:
            code_text.append(line)
            i += 1
            continue

        # Markdown Tables
        if line.startswith('|') and '|' in line[1:]:
            table_lines.append(line)
            i += 1
            continue
        else:
            if table_lines:
                flush_table(table_lines)
                table_lines = []

        if not line.strip():
            i += 1
            continue

        # Dividers
        if line.strip() == '---':
            i += 1
            continue

        # Headings
        if line.startswith('# '):
            h = doc.add_heading(level=0)
            run = h.add_run(line[2:].strip())
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0A, 0x25, 0x40)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            h.paragraph_format.space_after = Pt(10)
        elif line.startswith('## '):
            h = doc.add_heading(level=1)
            run = h.add_run(line[3:].strip())
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
        elif line.startswith('### '):
            h = doc.add_heading(level=2)
            run = h.add_run(line[4:].strip())
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
        elif line.startswith('#### '):
            h = doc.add_heading(level=3)
            run = h.add_run(line[5:].strip())
            run.font.size = Pt(10.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(3)
        elif line.startswith('- ') or line.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            content = line[2:].strip()
            # Simple bold formatting parser
            parts = content.split('**')
            for idx, part in enumerate(parts):
                r = p.add_run(part)
                if idx % 2 == 1:
                    r.font.bold = True
        elif line[0].isdigit() and (line[1:3] == '. ' or line[2:4] == '. '):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            dot_idx = line.find('. ')
            content = line[dot_idx+2:].strip()
            parts = content.split('**')
            for idx, part in enumerate(parts):
                r = p.add_run(part)
                if idx % 2 == 1:
                    r.font.bold = True
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            content = line.strip()
            parts = content.split('**')
            for idx, part in enumerate(parts):
                r = p.add_run(part)
                if idx % 2 == 1:
                    r.font.bold = True

        i += 1

    if table_lines:
        flush_table(table_lines)

    output_path = 'PROJECT_REPORT.docx'
    doc.save(output_path)
    print(f"Generated {output_path} successfully ({os.path.getsize(output_path)} bytes).")

if __name__ == '__main__':
    generate_report_docx()
